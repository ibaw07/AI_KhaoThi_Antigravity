"""
Module trích xuất văn bản từ các định dạng tài liệu DOCX, PDF, TXT.
Tối ưu hóa chạy cục bộ (100% Offline), không cần quyền Administrator.
"""

import os
from pathlib import Path
from typing import Union, Dict, Any


def extract_text_from_docx(file_path: Union[str, Path]) -> str:
    """
    Trích xuất toàn bộ văn bản từ file Word (.docx) bao gồm paragraphs và tables.
    """
    import docx

    doc = docx.Document(file_path)
    full_text = []

    # Đọc paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            full_text.append(text)

    # Đọc tables (giữ cấu trúc bảng nếu có)
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            # Loại bỏ ô trùng lặp do merge
            dedup_cells = []
            for c in row_cells:
                if not dedup_cells or c != dedup_cells[-1]:
                    dedup_cells.append(c)
            row_line = " | ".join(dedup_cells)
            if row_line.strip():
                table_rows.append(row_line)
        if table_rows:
            full_text.append("\n" + "\n".join(table_rows) + "\n")

    return "\n\n".join(full_text)


def extract_text_from_pdf(file_path: Union[str, Path]) -> str:
    """
    Trích xuất văn bản từ file PDF sử dụng pypdfium2 (nhanh, nhẹ, không cần Poppler).
    """
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(file_path)
    full_text = []

    for page_idx in range(len(pdf)):
        page = pdf[page_idx]
        textpage = page.get_textpage()
        text = textpage.get_text_range().strip()
        if text:
            full_text.append(text)

    return "\n\n".join(full_text)


def extract_text_from_txt(file_path: Union[str, Path]) -> str:
    """
    Trích xuất văn bản từ file TXT với các encoding phổ biến: utf-8, utf-8-sig, cp1258, latin-1.
    """
    encodings = ["utf-8", "utf-8-sig", "cp1258", "latin-1"]
    content = None

    with open(file_path, "rb") as f:
        raw_bytes = f.read()

    for enc in encodings:
        try:
            content = raw_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue

    if content is None:
        content = raw_bytes.decode("utf-8", errors="replace")

    return content.strip()


def extract_text_from_file(file_path: Union[str, Path]) -> Dict[str, Any]:
    """
    Hàm tổng quát trích xuất văn bản dựa trên phần mở rộng tệp.
    Trả về dict chứa:
    - filename: tên file
    - extension: đuôi file (.docx, .pdf, .txt)
    - text: nội dung văn bản trích xuất
    - word_count: số từ ước tính
    - char_count: số ký tự
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Không tìm thấy tệp: {file_path}")

    ext = path.suffix.lower()
    text = ""

    if ext == ".docx":
        text = extract_text_from_docx(path)
    elif ext == ".pdf":
        text = extract_text_from_pdf(path)
    elif ext in [".txt", ".md", ".log"]:
        text = extract_text_from_txt(path)
    else:
        raise ValueError(f"Định dạng tệp không được hỗ trợ: {ext}. Chỉ hỗ trợ .docx, .pdf, .txt")

    # Chuẩn hóa khoảng trắng dư thừa nhưng giữ dòng
    normalized_lines = []
    for line in text.splitlines():
        cleaned = " ".join(line.split())
        if cleaned:
            normalized_lines.append(cleaned)

    cleaned_text = "\n\n".join(normalized_lines)
    word_count = len(cleaned_text.split())

    return {
        "filename": path.name,
        "extension": ext,
        "text": cleaned_text,
        "word_count": word_count,
        "char_count": len(cleaned_text)
    }


# =========================================================================
# HÀM LỌC HẬU XỬ LÝ (POST-PROCESSING FILTER) CHO CÂU HỎI KHẢO THÍ
# =========================================================================

import re
from typing import List, Optional, Tuple

# Danh sách Regex nhận diện và cắt bỏ tiền tố siêu tham chiếu (meta-prefixes) ở đầu câu dẫn
META_PREFIX_PATTERNS = [
    r"^(Căn cứ vào|Dựa vào|Theo)\s+(dữ liệu\s+)?(văn bản|tài liệu|đoạn trích|bài đọc|đoạn văn)(\s+nguồn)?\s*([\'\"‘’“][^\'\"‘’“”]*[\'\"‘’”]|thuộc chủ đề\s+[\'\"‘’“][^\'\"‘’“”]*[\'\"‘’”]|trên|dưới đây|đã cho|nêu trên|được trích|đã nêu)?[:,\s]*",
    r"^(Trong|Ở)\s+(chủ đề|văn bản|tài liệu|đoạn trích|bài đọc)\s*([\'\"‘’“][^\'\"‘’“”]*[\'\"‘’”]|trên|dưới đây|đã cho|nêu trên)?[:,\s]*",
    r"^(Qua|Từ)\s+(nội dung\s+)?(văn bản|tài liệu|đoạn trích|bài đọc)\s*(trên|nêu trên|đã cho)?[:,\s]*",
    r"^Căn cứ vào\s+(các\s+)?(dữ liệu|thông tin|luận điểm)\s*(nêu trên|trên|đã cho)?[:,\s]*",
    r"^Theo\s+(tác giả|bài viết|bài đọc|đoạn văn)\s*(nêu trên|trên|đã cho)?[:,\s]*",
    r"^(Văn bản|Đoạn trích|Tài liệu)\s+(khẳng định|cho biết|đề cập|nêu rõ|xác định)\s+(rằng\s+)?[:,\s]*",
]

COMPILED_META_PREFIXES = [re.compile(p, re.IGNORECASE) for p in META_PREFIX_PATTERNS]

# Danh sách từ khóa rác cấm xuất hiện trong các phương án lựa chọn và câu dẫn
FORBIDDEN_META_WORDS = [
    "văn bản",
    "tài liệu",
    "ngoài tài liệu",
    "đoạn trích",
    "bài đọc"
]


def strip_meta_prefixes(stem: str) -> str:
    """
    Tự động cắt bỏ các tiền tố siêu tham chiếu ở đầu câu dẫn:
    Ví dụ: 'Căn cứ vào dữ liệu văn bản 'Chủ đề mở đầu', khẳng định nào sau đây là đúng?'
    -> 'Khẳng định nào sau đây là đúng?'
    """
    if not stem:
        return ""

    cleaned_stem = stem.strip()
    
    # Lặp qua các pattern để bóc tách triệt để nếu có tiền tố lồng nhau
    changed = True
    while changed:
        changed = False
        for pattern in COMPILED_META_PREFIXES:
            match = pattern.match(cleaned_stem)
            if match:
                cleaned_stem = cleaned_stem[match.end():].strip()
                changed = True
                break

    # Viết hoa chữ cái đầu tiên của câu dẫn sau khi cắt tiền tố
    if cleaned_stem:
        cleaned_stem = cleaned_stem[0].upper() + cleaned_stem[1:]

    return cleaned_stem


def extract_all_option_texts(options: Any) -> List[str]:
    """Trích xuất danh sách tất cả các chuỗi phương án từ options."""
    texts = []
    if isinstance(options, list):
        for opt in options:
            if opt is not None:
                texts.append(str(opt))
    elif isinstance(options, dict):
        for key in ["column_a", "column_b", "items", "choices"]:
            val = options.get(key)
            if isinstance(val, list):
                for item in val:
                    if item is not None:
                        texts.append(str(item))
    elif isinstance(options, str):
        texts.append(options)
    return texts


def validate_question_content(question: Dict[str, Any]) -> Tuple[bool, str]:
    """
    Kiểm tra câu hỏi và các phương án lựa chọn:
    Nếu phương án hoặc câu dẫn chứa các cụm từ rác ('văn bản', 'tài liệu', 'ngoài tài liệu'...),
    sẽ trả về (False, lý do) để tự động loại bỏ và yêu cầu tạo lại.
    """
    stem = str(question.get("stem") or "").lower()
    options = question.get("options")
    option_texts = extract_all_option_texts(options)

    # 1. Kiểm tra câu dẫn (sau khi đã chuẩn hóa)
    for forbidden in FORBIDDEN_META_WORDS:
        if forbidden in stem:
            return False, f"Câu dẫn chứa từ khóa cấm: '{forbidden}'"

    # 2. Kiểm tra tất cả các phương án A, B, C, D, E...
    for opt in option_texts:
        opt_lower = opt.lower()
        for forbidden in FORBIDDEN_META_WORDS:
            if forbidden in opt_lower:
                return False, f"Phương án lựa chọn chứa cụm từ rác: '{forbidden}' trong '{opt}'"

    return True, "Hợp lệ"


def post_process_filter_question(question: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    """
    Áp dụng bộ lọc hậu xử lý cho một câu hỏi:
    1. Cắt bỏ tiền tố siêu tham chiếu ở đầu câu dẫn (Stand-alone Question Stem).
    2. Kiểm tra tính hợp lệ của câu dẫn và phương án (loại bỏ phương án nhiễu rác).
    Trả về câu hỏi đã làm sạch nếu hợp lệ, hoặc None nếu vi phạm cần tạo lại.
    """
    if not question or not isinstance(question, dict):
        return None

    cleaned_q = dict(question)
    original_stem = str(cleaned_q.get("stem") or "")
    
    # 1. Cắt bỏ tiền tố
    cleaned_stem = strip_meta_prefixes(original_stem)
    if not cleaned_stem:
        return None
    cleaned_q["stem"] = cleaned_stem

    # 2. Kiểm định tính hợp lệ không chứa từ rác
    is_valid, reason = validate_question_content(cleaned_q)
    if not is_valid:
        return None

    return cleaned_q


def filter_and_clean_questions(questions: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """
    Lọc và làm sạch toàn bộ danh sách câu hỏi:
    Loại bỏ các câu hỏi chứa phương án rác hoặc câu dẫn vi phạm.
    """
    valid_questions = []
    for q in questions:
        cleaned = post_process_filter_question(q)
        if cleaned is not None:
            valid_questions.append(cleaned)
    return valid_questions

